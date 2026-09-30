import uuid
from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from django.contrib.auth.decorators import login_required
from django.views.decorators.http import require_POST
from django.contrib import messages
from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from django.db.models import Sum, Count, Exists, OuterRef, F, Q
from django.db.models.functions import Coalesce, TruncDate
from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.urls import reverse
from django.utils import timezone

from apps.accounts.decorators import platform_admin_required, inventory_access_required, inventory_manage_required, movements_access_required, reports_access_required
from apps.companies.models import Company, Branch
from .models import Category, Product, Stock, Movement

"""------------------------------------------------------------------ Category Views ------------------------------------------------------------------"""

@login_required
@platform_admin_required
def category_list(request):
    categories = Category.objects.order_by('name')
    return render(request, "categories/category_list.html", {'categories': categories})

@login_required
@platform_admin_required
def create_category(request):
    if request.method == "POST":
        name = request.POST.get("name", "").strip()

        base_context = {'name': name}

        errors = []

        if not name:
            errors.append("El nombre de la categoría es obligatorio.")
        if name and len(name) > 100:
            errors.append("El nombre de la categoría no puede superar los 100 caracteres.")
        if name and Category.objects.filter(name=name).exists():
            errors.append("Ya existe una categoría con ese nombre.")

        if errors:
            messages.error(request, errors[0])
            return render(request, 'categories/create_category.html', base_context)

        Category.objects.create(name=name)

        # Éxito: NO se hace redirect() directo. Se re-renderiza el form con `redirect_url`
        # para que el template muestre el SweetAlert y navegue recién al cerrarlo.
        messages.success(request, f"Categoría {name} creada correctamente.")
        return render(request, 'categories/create_category.html', {
            'redirect_url': reverse('inventory:category_list'),
        })

    # GET: form vacío.
    return render(request, 'categories/create_category.html')

@login_required
@platform_admin_required
def update_category(request, category_id):
    category = get_object_or_404(Category, id=category_id)
    context = {'category': category}

    if request.method == "POST":
        name = request.POST.get("name", "").strip()

        # Se re-renderiza el form conservando lo que el admin ya había escrito, si hay que volver a mostrarlo.
        submitted_context = {**context, 'name': name}

        errors = []

        if not name:
            errors.append("El nombre de la categoría es obligatorio.")
        if name and len(name) > 100:
            errors.append("El nombre de la categoría no puede superar los 100 caracteres.")
        # exclude(id=category.id): que la propia categoría no cuente como "duplicada" de sí misma.
        if name and Category.objects.filter(name=name).exclude(id=category.id).exists():
            errors.append("Ya existe una categoría con ese nombre.")

        if errors:
            messages.error(request, errors[0])
            return render(request, 'categories/update_category.html', submitted_context)

        try:
            category.name = name
            category.full_clean()
            category.save()
        except (IntegrityError, ValidationError):
            messages.error(request, "Ocurrió un error al actualizar la categoría. Inténtalo de nuevo.")
            return render(request, "categories/update_category.html", submitted_context)

        # Éxito: NO se hace redirect() directo. Se re-renderiza el form con `redirect_url`
        # para que el template muestre el SweetAlert y navegue recién al cerrarlo.
        messages.success(request, f"La categoría {name} fue editada exitosamente.")
        context['redirect_url'] = reverse('inventory:category_list')

    return render(request, 'categories/update_category.html', context)

# Elimina una categoría. SET_NULL: los productos que la tenían quedan sin categoría (Product.category on_delete=SET_NULL).
@login_required
@require_POST
@platform_admin_required
def delete_category(request, category_id):
    category = get_object_or_404(Category, id=category_id)
    category.delete()
    return redirect('inventory:category_list')

"""------------------------------------------------------------------ Product Views ------------------------------------------------------------------"""

# Resuelve sobre qué empresa se trabaja: la propia del usuario, o -si es admin de plataforma
# dando soporte- la que eligió por company_id. None si no hay una empresa válida para ver.
def _resolve_company(request, company_id):
    if company_id is not None:
        if not request.user.is_platform_admin:
            return None
        return get_object_or_404(Company, id=company_id)
    if request.user.is_platform_admin:
        return None
    return request.user.company

def _product_list_url(request, company):
    if request.user.is_platform_admin:
        return reverse('inventory:product_list', args=[company.id])
    return reverse('inventory:product_list')

# Sucursales sobre las que puede operar quien hace el pedido: todas si es admin de plataforma
# (soporte) o un usuario "de empresa" (User.branch=None); solo la suya si tiene una sucursal
# asignada. Sin esto, un empleado de una sucursal podía ver/mover stock de las demás.
def _allowed_branches(request, company):
    branches_qs = company.branches.filter(is_active=True).order_by('name')
    if not request.user.is_platform_admin and request.user.branch_id:
        return branches_qs.filter(id=request.user.branch_id)
    return branches_qs

@login_required
@inventory_access_required
def product_list(request, company_id=None):
    company = _resolve_company(request, company_id)
    if company is None:
        if request.user.is_platform_admin:
            messages.error(request, "Selecciona una empresa desde el listado.")
            return redirect('company:list')
        messages.error(request, "No tienes permiso para ver esta empresa.")
        return redirect('dashboard')

    # Solo se ve/suma el stock de las sucursales permitidas (ver _allowed_branches): un usuario
    # con una sola sucursal asignada no debe poder inferir cuánto hay en las demás.
    allowed_branches = _allowed_branches(request, company)
    # "Bajo" es por sucursal, no por el total sumado: una sucursal en 0 con otra sobrada
    # da un total que parece sano, pero esa sucursal puntual sí necesita reabastecerse.
    low_stock_subquery = Stock.objects.filter(product=OuterRef('pk'), branch__in=allowed_branches, quantity__lte=F('product__min_stock'))
    products = (
        Product.objects.filter(company=company)
        .select_related('category')
        .annotate(
            total_stock=Coalesce(Sum('stocks__quantity', filter=Q(stocks__branch__in=allowed_branches)), 0),
            has_low_stock=Exists(low_stock_subquery),
        )
        .order_by('name')
    )
    return render(request, "products/product_list.html", {'company': company, 'products': products})

@login_required
@inventory_access_required
def product_detail(request, product_id):
    product = get_object_or_404(Product.objects.select_related('company', 'category'), id=product_id)

    if not request.user.is_platform_admin and product.company_id != request.user.company_id:
        messages.error(request, "No tienes permiso para ver este producto.")
        return redirect('dashboard')

    allowed_branches = _allowed_branches(request, product.company)
    stocks = Stock.objects.filter(product=product, branch__in=allowed_branches).select_related('branch').order_by('branch__name')
    total_stock = sum(s.quantity for s in stocks)
    return render(request, "products/product_detail.html", {
        'product': product, 'stocks': stocks, 'total_stock': total_stock,
    })

@login_required
@inventory_manage_required
def create_product(request, company_id=None):
    company = _resolve_company(request, company_id)
    if company is None:
        if request.user.is_platform_admin:
            messages.error(request, "Selecciona una empresa desde el listado.")
            return redirect('company:list')
        messages.error(request, "No tienes permiso para ver esta empresa.")
        return redirect('dashboard')

    categories_qs = Category.objects.order_by('name')
    # El front usa esto para armar un SKU que no choque con uno ya usado por esta empresa
    # (ver static/js/products/product_sku.script.js), sin tener que ir preguntándole al server.
    existing_skus = list(Product.objects.filter(company=company).values_list('sku', flat=True))

    if request.method == "POST":
        name = request.POST.get("name", "").strip()
        sku = request.POST.get("sku", "").strip()
        unit = request.POST.get("unit", "").strip()
        category_id = request.POST.get("category", "").strip()
        cost_price_raw = request.POST.get("cost_price", "").strip()
        sale_price_raw = request.POST.get("sale_price", "").strip()
        min_stock_raw = request.POST.get("min_stock", "").strip()

        base_context = {
            'company': company,
            'categories': categories_qs,
            'existing_skus': existing_skus,
            'name': name, 'sku': sku, 'unit': unit, 'selected_category': category_id,
            'cost_price': cost_price_raw, 'sale_price': sale_price_raw, 'min_stock': min_stock_raw,
        }

        errors = []

        if not name:
            errors.append("El nombre del producto es obligatorio.")
        if not sku:
            errors.append("El código (SKU) es obligatorio.")
        if sku and Product.objects.filter(company=company, sku=sku).exists():
            errors.append("Ya existe un producto con ese código en esta empresa.")
        if not category_id:
            errors.append("Selecciona una categoría.")
        if not unit:
            errors.append("La unidad es obligatoria.")

        cost_price = sale_price = None
        if not cost_price_raw:
            errors.append("El precio de costo es obligatorio.")
        else:
            try:
                cost_price = Decimal(cost_price_raw)
                if cost_price <= 0:
                    errors.append("El precio de costo debe ser mayor a cero.")
            except InvalidOperation:
                errors.append("El precio de costo no es válido.")

        if not sale_price_raw:
            errors.append("El precio de venta es obligatorio.")
        else:
            try:
                sale_price = Decimal(sale_price_raw)
                if sale_price <= 0:
                    errors.append("El precio de venta debe ser mayor a cero.")
            except InvalidOperation:
                errors.append("El precio de venta no es válido.")

        if cost_price is not None and sale_price is not None and sale_price <= cost_price:
            errors.append("El precio de venta debe ser mayor al precio de costo.")

        min_stock = None
        if not min_stock_raw:
            errors.append("El stock mínimo es obligatorio.")
        elif not min_stock_raw.isdigit():
            errors.append("El stock mínimo debe ser un número entero.")
        elif int(min_stock_raw) < 5:
            errors.append("El stock mínimo no puede ser menor a 5.")
        else:
            min_stock = int(min_stock_raw)

        if errors:
            messages.error(request, errors[0])
            return render(request, 'products/create_product.html', base_context)

        category_obj = get_object_or_404(Category, id=category_id)

        # Stock inicial: 10 unidades en cada sucursal activa, cada una como su propio Movement
        # (tipo Entrada, motivo "ajuste") en vez de un Stock suelto, para no romper la regla de
        # que el stock SIEMPRE cambia a través de un movimiento auditable.
        INITIAL_STOCK_QTY = 10
        try:
            with transaction.atomic():
                product = Product(
                    company=company, category=category_obj,
                    name=name, sku=sku, unit=unit,
                    cost_price=cost_price, sale_price=sale_price, min_stock=min_stock,
                )
                product.full_clean()
                product.save()

                for branch in company.branches.filter(is_active=True):
                    Stock.objects.create(product=product, branch=branch, quantity=INITIAL_STOCK_QTY)
                    movement = Movement(
                        company=company, branch=branch, product=product, user=request.user,
                        movement_type=Movement.IN, reason='ajuste', quantity=INITIAL_STOCK_QTY,
                        unit_price=product.cost_price, note="Stock inicial al crear el producto.",
                    )
                    movement.full_clean()
                    movement.save()
        except (IntegrityError, ValidationError):
            messages.error(request, "Ocurrió un error al crear el producto. Inténtalo de nuevo.")
            return render(request, 'products/create_product.html', base_context)

        # Éxito: NO se hace redirect() directo. Se re-renderiza el form con `redirect_url`
        # para que el template muestre el SweetAlert y navegue recién al cerrarlo.
        messages.success(request, f"Producto {name} creado correctamente, con {INITIAL_STOCK_QTY} unidades de stock inicial en cada sucursal.")

        return render(request, 'products/create_product.html', {
            'company': company,
            'categories': categories_qs,
            'existing_skus': existing_skus + [sku],
            'redirect_url': _product_list_url(request, company),
        })

    # GET: form vacío.
    return render(request, 'products/create_product.html', {
        'company': company, 'categories': categories_qs, 'existing_skus': existing_skus,
    })

@login_required
@inventory_manage_required
def update_product(request, product_id):
    product = get_object_or_404(Product, id=product_id)

    if not request.user.is_platform_admin and product.company_id != request.user.company_id:
        messages.error(request, "No tienes permiso para editar este producto.")
        return redirect('dashboard')

    categories_qs = Category.objects.order_by('name')
    context = {'product': product, 'categories': categories_qs}

    if request.method == "POST":
        name = request.POST.get("name", "").strip()
        sku = request.POST.get("sku", "").strip()
        unit = request.POST.get("unit", "").strip()
        category_id = request.POST.get("category", "").strip()
        cost_price_raw = request.POST.get("cost_price", "").strip()
        sale_price_raw = request.POST.get("sale_price", "").strip()
        min_stock_raw = request.POST.get("min_stock", "").strip()

        # Se re-renderiza el form conservando lo que ya se había escrito, si hay que volver a mostrarlo.
        submitted_context = {
            **context,
            'name': name, 'sku': sku, 'unit': unit, 'selected_category': category_id,
            'cost_price': cost_price_raw, 'sale_price': sale_price_raw, 'min_stock': min_stock_raw,
        }

        errors = []

        if not name:
            errors.append("El nombre del producto es obligatorio.")
        if not sku:
            errors.append("El código (SKU) es obligatorio.")
        # exclude(id=product.id): que el propio producto no cuente como "duplicado" de sí mismo.
        if sku and Product.objects.filter(company=product.company, sku=sku).exclude(id=product.id).exists():
            errors.append("Ya existe un producto con ese código en esta empresa.")
        if not category_id:
            errors.append("Selecciona una categoría.")
        if not unit:
            errors.append("La unidad es obligatoria.")

        cost_price = sale_price = None
        if not cost_price_raw:
            errors.append("El precio de costo es obligatorio.")
        else:
            try:
                cost_price = Decimal(cost_price_raw)
                if cost_price <= 0:
                    errors.append("El precio de costo debe ser mayor a cero.")
            except InvalidOperation:
                errors.append("El precio de costo no es válido.")

        if not sale_price_raw:
            errors.append("El precio de venta es obligatorio.")
        else:
            try:
                sale_price = Decimal(sale_price_raw)
                if sale_price <= 0:
                    errors.append("El precio de venta debe ser mayor a cero.")
            except InvalidOperation:
                errors.append("El precio de venta no es válido.")

        if cost_price is not None and sale_price is not None and sale_price <= cost_price:
            errors.append("El precio de venta debe ser mayor al precio de costo.")

        min_stock = None
        if not min_stock_raw:
            errors.append("El stock mínimo es obligatorio.")
        elif not min_stock_raw.isdigit():
            errors.append("El stock mínimo debe ser un número entero.")
        elif int(min_stock_raw) < 5:
            errors.append("El stock mínimo no puede ser menor a 5.")
        else:
            min_stock = int(min_stock_raw)

        if errors:
            messages.error(request, errors[0])
            return render(request, 'products/update_product.html', submitted_context)

        category_obj = get_object_or_404(Category, id=category_id)

        try:
            product.name = name
            product.sku = sku
            product.unit = unit
            product.category = category_obj
            product.cost_price = cost_price
            product.sale_price = sale_price
            product.min_stock = min_stock
            product.full_clean()
            product.save()
        except (IntegrityError, ValidationError):
            messages.error(request, "Ocurrió un error al actualizar el producto. Inténtalo de nuevo.")
            return render(request, "products/update_product.html", submitted_context)

        # Éxito: NO se hace redirect() directo. Se re-renderiza el form con `redirect_url`
        # para que el template muestre el SweetAlert y navegue recién al cerrarlo.
        messages.success(request, f"El producto {name} fue editado exitosamente.")

        submitted_context['redirect_url'] = _product_list_url(request, product.company)
        return render(request, 'products/update_product.html', submitted_context)

    return render(request, 'products/update_product.html', context)

# Baja lógica (is_active=False): un producto no se borra nunca, tiene Movement/Stock asociados
# que forman el historial de inventario.
@login_required
@require_POST
@inventory_manage_required
def deactivate_product(request, product_id):
    product = get_object_or_404(Product, id=product_id)
    if not request.user.is_platform_admin and product.company_id != request.user.company_id:
        messages.error(request, "No tienes permiso para editar este producto.")
        return redirect('dashboard')
    product.is_active = False
    product.save()
    messages.success(request, f"{product.name} fue desactivado correctamente.")
    return redirect(_product_list_url(request, product.company))

@login_required
@require_POST
@inventory_manage_required
def activate_product(request, product_id):
    product = get_object_or_404(Product, id=product_id)
    if not request.user.is_platform_admin and product.company_id != request.user.company_id:
        messages.error(request, "No tienes permiso para editar este producto.")
        return redirect('dashboard')
    product.is_active = True
    product.save()
    messages.success(request, f"{product.name} fue reactivado correctamente.")
    return redirect(_product_list_url(request, product.company))

"""------------------------------------------------------------------ Movement Views ------------------------------------------------------------------"""

def _movement_list_url(request, company):
    if request.user.is_platform_admin:
        return reverse('inventory:movement_list', args=[company.id])
    return reverse('inventory:movement_list')

@login_required
@movements_access_required
def movement_list(request, company_id=None):
    company = _resolve_company(request, company_id)
    if company is None:
        if request.user.is_platform_admin:
            messages.error(request, "Selecciona una empresa desde el listado.")
            return redirect('company:list')
        messages.error(request, "No tienes permiso para ver esta empresa.")
        return redirect('dashboard')

    movements = (
        Movement.objects.filter(company=company, branch__in=_allowed_branches(request, company))
        .select_related('product', 'branch', 'user')
        .order_by('-created_at')
    )
    return render(request, "movements/movement_list.html", {'company': company, 'movements': movements})

@login_required
@movements_access_required
def create_movement(request, company_id=None):
    company = _resolve_company(request, company_id)
    if company is None:
        if request.user.is_platform_admin:
            messages.error(request, "Selecciona una empresa desde el listado.")
            return redirect('company:list')
        messages.error(request, "No tienes permiso para ver esta empresa.")
        return redirect('dashboard')

    branches_qs = _allowed_branches(request, company)
    # Solo productos con Stock ya registrado en alguna de las sucursales permitidas: no tiene
    # sentido ofrecer en el selector productos que esa sucursal nunca tuvo (pedido explícito
    # del negocio, aunque implica que un producto recién creado necesita su primer Stock por
    # otra vía antes de poder registrarle un movimiento).
    products_qs = (
        Product.objects.filter(company=company, is_active=True)
        .filter(Exists(Stock.objects.filter(product=OuterRef('pk'), branch__in=branches_qs)))
        .order_by('name')
    )

    if request.method == "POST":
        branch_id = request.POST.get("branch", "").strip()
        movement_type = request.POST.get("movement_type", "").strip()
        reason = request.POST.get("reason", "").strip()
        note = request.POST.get("note", "").strip()
        # Varias filas Producto+Cantidad en paralelo (un solo movimiento "de compra" puede
        # traer varios productos a la vez, como una factura), comparten sucursal/tipo/motivo/nota.
        product_ids_posted = request.POST.getlist("product[]")
        quantities_posted = request.POST.getlist("quantity[]")
        posted_rows = [{'product': pid, 'quantity': qty} for pid, qty in zip(product_ids_posted, quantities_posted)]

        base_context = {
            'company': company, 'products': products_qs, 'branches': branches_qs,
            'selected_branch': branch_id, 'selected_movement_type': movement_type,
            'selected_reason': reason, 'note': note,
            'product_rows': posted_rows or [{'product': '', 'quantity': ''}],
        }

        errors = []

        if not branch_id:
            errors.append("Selecciona la sucursal.")
        if movement_type not in (Movement.IN, Movement.OUT):
            errors.append("Selecciona el tipo de movimiento.")
        if not reason:
            errors.append("Selecciona el motivo.")
        elif movement_type in Movement.REASONS_BY_TYPE and reason not in Movement.REASONS_BY_TYPE[movement_type]:
            errors.append("Ese motivo no aplica para este tipo de movimiento.")

        # Filas totalmente vacías (fila de más que el usuario no llegó a usar) se ignoran.
        filled_rows = [
            (pid.strip(), qty.strip())
            for pid, qty in zip(product_ids_posted, quantities_posted)
            if pid.strip() or qty.strip()
        ]

        if not filled_rows:
            errors.append("Agrega al menos un producto.")

        parsed_rows = []  # [(product_id, quantity), ...]
        seen_product_ids = set()
        for pid, qty_raw in filled_rows:
            if not pid:
                errors.append("Selecciona el producto en cada fila.")
                break
            if not qty_raw or not qty_raw.isdigit() or int(qty_raw) <= 0:
                errors.append("La cantidad de cada producto debe ser un número entero mayor a cero.")
                break
            if pid in seen_product_ids:
                errors.append("No repitas el mismo producto en dos filas: sumá la cantidad en una sola.")
                break
            seen_product_ids.add(pid)
            parsed_rows.append((pid, int(qty_raw)))

        if errors:
            messages.error(request, errors[0])
            return render(request, 'movements/create_movement.html', base_context)

        # branch_obj sale de branches_qs (ya filtrado por _allowed_branches), no de un
        # get_object_or_404(company=company) suelto: así alguien con una sola sucursal asignada
        # no puede mandar a mano el id de otra sucursal de la misma empresa y saltarse el límite.
        branch_obj = get_object_or_404(branches_qs, id=branch_id)

        # Un solo batch_id para todas las filas de este envío (aunque sea una sola): así el
        # listado las puede agrupar como "un mismo movimiento" en vez de registros sueltos.
        batch_id = uuid.uuid4()

        try:
            with transaction.atomic():
                results = []  # [(product_obj, resulting_quantity), ...]
                for pid, quantity in parsed_rows:
                    product_obj = get_object_or_404(Product, id=pid, company=company)
                    stock_obj, _ = Stock.objects.get_or_create(product=product_obj, branch=branch_obj, defaults={'quantity': 0})

                    if movement_type == Movement.OUT and quantity > stock_obj.quantity:
                        raise ValidationError(f"No hay suficiente stock de {product_obj.name} en {branch_obj.name} (actual: {stock_obj.quantity}).")

                    stock_obj.quantity += quantity if movement_type == Movement.IN else -quantity
                    stock_obj.save()

                    unit_price = product_obj.cost_price if movement_type == Movement.IN else product_obj.sale_price
                    movement = Movement(
                        company=company, branch=branch_obj, product=product_obj, user=request.user,
                        movement_type=movement_type, reason=reason, quantity=quantity,
                        unit_price=unit_price, batch_id=batch_id, note=note,
                    )
                    movement.full_clean()
                    movement.save()
                    results.append((product_obj, stock_obj.quantity))
        except ValidationError as e:
            messages.error(request, e.messages[0])
            return render(request, 'movements/create_movement.html', base_context)
        except IntegrityError:
            messages.error(request, "Ocurrió un error al registrar el movimiento. Inténtalo de nuevo.")
            return render(request, 'movements/create_movement.html', base_context)

        # Éxito: NO se hace redirect() directo. Se re-renderiza el form con `redirect_url`
        # para que el template muestre el SweetAlert y navegue recién al cerrarlo.
        messages.success(request, f"Movimiento registrado: {len(results)} producto(s) en {branch_obj.name}.")

        if movement_type == Movement.OUT:
            low_products = [p.name for p, qty in results if qty <= p.min_stock]
            if low_products:
                messages.warning(request, f"{branch_obj.name} quedó con stock bajo en: {', '.join(low_products)}.")

        return render(request, 'movements/create_movement.html', {
            'company': company, 'products': products_qs, 'branches': branches_qs,
            'product_rows': [{'product': '', 'quantity': ''}],
            'redirect_url': _movement_list_url(request, company),
        })

    # GET: form vacío.
    return render(request, 'movements/create_movement.html', {
        'company': company, 'products': products_qs, 'branches': branches_qs,
        'product_rows': [{'product': '', 'quantity': ''}],
    })

"""------------------------------------------------------------------ Report Views ------------------------------------------------------------------"""

def _report_url(request, company):
    if request.user.is_platform_admin:
        return reverse('inventory:report', args=[company.id])
    return reverse('inventory:report')

# Excel con diseño real (encabezado con color, negrita, columnas ajustadas al contenido, filas
# alternadas) en vez de un CSV a secas. Recibe el queryset YA filtrado por report_view.
def _movements_to_xlsx(movements, date_from, date_to):
    wb = Workbook()
    ws = wb.active
    ws.title = "Movimientos"

    headers = ['Producto', 'Sucursal', 'Tipo', 'Cantidad', 'Motivo', 'Precio unitario', 'Valor total', 'Usuario', 'Fecha']
    ws.append(headers)

    header_fill = PatternFill(start_color="FF1657D6", end_color="FF1657D6", fill_type="solid")
    header_font = Font(color="FFFFFFFF", bold=True)
    thin_border = Border(bottom=Side(style='thin', color='FFDDDDDD'))
    stripe_fill = PatternFill(start_color="FFF2F5FB", end_color="FFF2F5FB", fill_type="solid")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    in_fill = PatternFill(start_color="FFDFF3E3", end_color="FFDFF3E3", fill_type="solid")
    out_fill = PatternFill(start_color="FFFCE1E7", end_color="FFFCE1E7", fill_type="solid")

    money_format = '#,##0'
    total_value = 0
    for i, m in enumerate(movements, start=2):
        ws.append([
            m.product.name, m.branch.name, m.get_movement_type_display(),
            m.quantity, m.get_reason_display(), float(m.unit_price), float(m.total_value),
            m.user.full_name if m.user else '—', m.created_at.strftime('%d/%m/%Y %H:%M'),
        ])
        total_value += m.total_value if m.movement_type == Movement.IN else -m.total_value
        row_fill = stripe_fill if i % 2 == 0 else None
        type_fill = in_fill if m.movement_type == Movement.IN else out_fill
        for col, cell in enumerate(ws[i], start=1):
            cell.border = thin_border
            cell.fill = type_fill if col == 3 else (row_fill or PatternFill())
            if col in (6, 7):
                cell.number_format = money_format

    # Fila final: valor neto (entradas - salidas) del rango filtrado, para que el contador no
    # tenga que sumar manualmente en Excel.
    total_row = ws.max_row + 1
    ws.cell(row=total_row, column=5, value="Neto (entradas - salidas):").font = Font(bold=True)
    net_cell = ws.cell(row=total_row, column=7, value=float(total_value))
    net_cell.font = Font(bold=True)
    net_cell.number_format = money_format

    for col_idx, header in enumerate(headers, start=1):
        max_len = max([len(header)] + [len(str(ws.cell(row=r, column=col_idx).value or '')) for r in range(2, ws.max_row + 1)])
        ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 4

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="movimientos_{date_from}_a_{date_to}.xlsx"'
    wb.save(response)
    return response

# Reportes es de SOLO LECTURA: reutiliza el mismo alcance de sucursales/movimientos que
# Movimientos, pero sin permitir crear/editar nada (can_access_reports es un permiso aparte).
@login_required
@reports_access_required
def report_view(request, company_id=None):
    company = _resolve_company(request, company_id)
    if company is None:
        if request.user.is_platform_admin:
            messages.error(request, "Selecciona una empresa desde el listado.")
            return redirect('company:list')
        messages.error(request, "No tienes permiso para ver esta empresa.")
        return redirect('dashboard')

    branches_qs = _allowed_branches(request, company)
    products_qs = Product.objects.filter(company=company).order_by('name')

    today = timezone.localdate()
    default_from = today - timedelta(days=30)

    date_from_raw = request.GET.get('date_from', '').strip()
    date_to_raw = request.GET.get('date_to', '').strip()
    branch_filter = request.GET.get('branch', '').strip()
    product_filter = request.GET.get('product', '').strip()
    type_filter = request.GET.get('movement_type', '').strip()

    try:
        date_from = datetime.strptime(date_from_raw, '%Y-%m-%d').date() if date_from_raw else default_from
    except ValueError:
        date_from = default_from
    try:
        date_to = datetime.strptime(date_to_raw, '%Y-%m-%d').date() if date_to_raw else today
    except ValueError:
        date_to = today

    movements = (
        Movement.objects.filter(
            company=company, branch__in=branches_qs,
            created_at__date__gte=date_from, created_at__date__lte=date_to,
        )
        .select_related('product', 'branch', 'user')
    )
    if branch_filter:
        movements = movements.filter(branch_id=branch_filter)
    if product_filter:
        movements = movements.filter(product_id=product_filter)
    if type_filter in (Movement.IN, Movement.OUT):
        movements = movements.filter(movement_type=type_filter)
    movements = movements.order_by('-created_at')

    # Exportar Excel: mismos filtros de arriba, se corta acá antes de armar gráficos/KPIs
    # (no hace falta calcular nada de eso para descargar un archivo).
    if request.GET.get('export') == 'xlsx':
        return _movements_to_xlsx(movements, date_from, date_to)

    total_in = sum(m.quantity for m in movements if m.movement_type == Movement.IN)
    total_out = sum(m.quantity for m in movements if m.movement_type == Movement.OUT)
    # Control monetario (para Contador/Gerencia): plata que entró (compras/costos) vs. plata
    # que salió/se facturó (ventas), a precio congelado en cada movimiento (Movement.unit_price).
    value_in = sum(m.total_value for m in movements if m.movement_type == Movement.IN)
    value_out = sum(m.total_value for m in movements if m.movement_type == Movement.OUT)

    low_stock_count = Stock.objects.filter(
        branch__in=branches_qs, quantity__lte=F('product__min_stock'), product__is_active=True,
    ).values('product_id').distinct().count()

    # Serie diaria (entradas vs salidas) para el gráfico de línea.
    daily = (
        movements
        .annotate(day=TruncDate('created_at'))
        .values('day', 'movement_type')
        .annotate(total=Sum('quantity'))
        .order_by('day')
    )
    daily_map = {}
    for row in daily:
        daily_map.setdefault(row['day'], {'IN': 0, 'OUT': 0})[row['movement_type']] = row['total']

    chart_labels = []
    chart_in = []
    chart_out = []
    cursor = date_from
    while cursor <= date_to:
        chart_labels.append(cursor.strftime('%d/%m'))
        day_totals = daily_map.get(cursor, {'IN': 0, 'OUT': 0})
        chart_in.append(day_totals.get('IN', 0))
        chart_out.append(day_totals.get('OUT', 0))
        cursor += timedelta(days=1)

    # Top 5 productos más movidos (por cantidad) en el rango filtrado.
    top_products = (
        movements.values('product__name')
        .annotate(total=Sum('quantity'))
        .order_by('-total')[:5]
    )

    # Stock bajo por sucursal (sobre las sucursales permitidas, sin importar el rango de fechas).
    low_stock_by_branch = (
        Stock.objects.filter(branch__in=branches_qs, quantity__lte=F('product__min_stock'), product__is_active=True)
        .values('branch__name')
        .annotate(total=Count('id'))
        .order_by('-total')
    )

    return render(request, 'reports/report_view.html', {
        'company': company, 'branches': branches_qs, 'products': products_qs,
        'movements': movements,
        'date_from': date_from.strftime('%Y-%m-%d'), 'date_to': date_to.strftime('%Y-%m-%d'),
        'selected_branch': branch_filter, 'selected_product': product_filter, 'selected_type': type_filter,
        'total_in': total_in, 'total_out': total_out, 'low_stock_count': low_stock_count,
        'value_in': value_in, 'value_out': value_out,
        'chart_labels': chart_labels, 'chart_in': chart_in, 'chart_out': chart_out,
        'top_products_labels': [p['product__name'] for p in top_products],
        'top_products_values': [p['total'] for p in top_products],
        'low_stock_labels': [b['branch__name'] for b in low_stock_by_branch],
        'low_stock_values': [b['total'] for b in low_stock_by_branch],
    })
